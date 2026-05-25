from __future__ import annotations

import csv
from dataclasses import asdict
from pathlib import Path


from app.core.spot_screener import SpotScanRow


def export_spot_scan_csv(rows: list[SpotScanRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = [asdict(r) for r in rows]
    headers = list(payload[0].keys()) if payload else [f.name for f in SpotScanRow.__dataclass_fields__.values()]
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        for row in payload:
            writer.writerow(row)
    return path


def export_spot_scan_xlsx(rows: list[SpotScanRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    headers = [f for f in SpotScanRow.__dataclass_fields__.keys()]
    ws.append(headers)
    fill_long = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_short = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_up = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    fill_down = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    for row in rows:
        vals = [getattr(row, h) for h in headers]
        ws.append(vals)
        excel_row = ws.max_row
        fill = None
        if row.long_signal:
            fill = fill_long
        elif row.short_signal:
            fill = fill_short
        elif row.trend_up:
            fill = fill_up
        elif row.trend_down:
            fill = fill_down
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=excel_row, column=col).fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path
